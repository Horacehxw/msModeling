"""Tests for output formatter modules."""

import pytest
from pathlib import Path
import tempfile

from tensor_cast.scripts.profiling_comparison.output import (
    BaseFormatter,
    ComparisonResult,
    ExcelFormatter,
    FormatterProtocol,
)
from tensor_cast.scripts.profiling_comparison.output.base import (
    ComparisonSummary,
    OperationMatch,
)


class TestOperationMatch:
    """Tests for OperationMatch dataclass."""

    def test_basic_creation(self):
        """Test basic OperationMatch creation."""
        match = OperationMatch(
            vllm_op="MatMulV2",
            tc_ops=["aten.mm"],
            vllm_duration_us=1000.0,
            tc_duration_us=950.0,
            vllm_count=10,
            tc_count=10,
            difference_us=-50.0,
            difference_pct=-5.0,
            match_status="exact",
        )
        assert match.vllm_op == "MatMulV2"
        assert match.tc_ops == ["aten.mm"]
        assert match.difference_pct == -5.0
        assert match.match_status == "exact"

    def test_missing_in_tc(self):
        """Test match with missing TC operation."""
        match = OperationMatch(
            vllm_op="CustomOp",
            tc_ops=[],
            vllm_duration_us=500.0,
            tc_duration_us=0.0,
            match_status="missing_in_tc",
        )
        assert match.tc_ops == []
        assert match.match_status == "missing_in_tc"


class TestComparisonSummary:
    """Tests for ComparisonSummary dataclass."""

    def test_basic_creation(self):
        """Test basic ComparisonSummary creation."""
        summary = ComparisonSummary(
            vllm_total_time_us=10000.0,
            tc_total_time_us=9500.0,
            overall_diff_us=-500.0,
            overall_diff_pct=-5.0,
            num_matched=20,
            num_missing_in_tc=2,
            num_missing_in_vllm=1,
            coverage_pct=95.0,
            avg_abs_diff_pct=8.5,
        )
        assert summary.vllm_total_time_us == 10000.0
        assert summary.num_matched == 20
        assert summary.coverage_pct == 95.0


class TestComparisonResult:
    """Tests for ComparisonResult dataclass."""

    def _create_sample_result(self):
        """Create a sample ComparisonResult for testing."""
        matches = [
            OperationMatch("Op1", ["tc1"], 1000, 950, 10, 10, -50, -5.0, "exact"),
            OperationMatch("Op2", ["tc2"], 500, 600, 5, 5, 100, 20.0, "partial"),
            OperationMatch("Op3", [], 200, 0, 2, 0, -200, -100, "missing_in_tc"),
        ]
        summary = ComparisonSummary(
            1700, 1550, -150, -8.8, 2, 1, 0, 88.2, 12.5
        )
        return ComparisonResult(
            config={"model_id": "test"},
            vllm_operations=[{"op_type": "Op1"}],
            tc_operations=[{"op_name": "tc1"}],
            matches=matches,
            summary=summary,
        )

    def test_get_matches_by_status(self):
        """Test filtering matches by status."""
        result = self._create_sample_result()

        exact_matches = result.get_matches_by_status("exact")
        assert len(exact_matches) == 1
        assert exact_matches[0].vllm_op == "Op1"

        missing = result.get_matches_by_status("missing_in_tc")
        assert len(missing) == 1

    def test_get_top_differences(self):
        """Test getting top differences."""
        result = self._create_sample_result()

        top = result.get_top_differences(2)
        assert len(top) == 2
        # Should be sorted by absolute difference
        assert abs(top[0].difference_pct) >= abs(top[1].difference_pct)


class TestExcelFormatter:
    """Tests for ExcelFormatter class."""

    def _create_sample_result(self):
        """Create a sample ComparisonResult for testing."""
        matches = [
            OperationMatch("MatMulV2", ["aten.mm"], 1000, 950, 10, 10, -50, -5.0, "exact"),
            OperationMatch("AddRmsNorm", ["aten.add", "tensor_cast.rmsnorm"], 500, 600, 5, 5, 100, 20.0, "partial"),
        ]
        summary = ComparisonSummary(
            1500, 1550, 50, 3.3, 2, 0, 0, 100.0, 12.5
        )
        return ComparisonResult(
            config={
                "model_id": "Qwen/Qwen3-32B",
                "device": "TEST",
                "world_size": 16,
                "tp_size": 16,
                "dp_size": 1,
                "ep": False,
                "quantize_linear_action": "DISABLED",
                "num_queries": 136,
                "query_length": 1,
                "context_length": 4096,
            },
            vllm_operations=[
                {
                    "op_type": "MatMulV2",
                    "count": 10,
                    "total_duration_us": 1000,
                    "avg_duration_us": 100,
                    "core_types": ["AI_CORE"],
                    "input_shapes": ["136,4096"],
                },
                {
                    "op_type": "AddRmsNorm",
                    "count": 5,
                    "total_duration_us": 500,
                    "avg_duration_us": 100,
                    "core_types": ["AI_CORE"],
                    "input_shapes": [],
                },
            ],
            tc_operations=[
                {
                    "op_name": "aten.mm",
                    "count": 10,
                    "total_time_us": 950,
                    "avg_time_us": 95,
                    "bound_classification": "compute",
                    "input_shapes": ["[136, 4096]"],
                },
                {
                    "op_name": "aten.add",
                    "count": 5,
                    "total_time_us": 250,
                    "avg_time_us": 50,
                    "bound_classification": "memory",
                    "input_shapes": [],
                },
                {
                    "op_name": "tensor_cast.rmsnorm",
                    "count": 5,
                    "total_time_us": 350,
                    "avg_time_us": 70,
                    "bound_classification": "memory",
                    "input_shapes": [],
                },
            ],
            matches=matches,
            summary=summary,
        )

    def test_format_creates_file(self):
        """Test that format creates an Excel file."""
        formatter = ExcelFormatter()
        result = self._create_sample_result()

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            output_path = Path(f.name)

        formatter.format(result, output_path)

        assert output_path.exists()
        # Check file has content
        assert output_path.stat().st_size > 0

    def test_format_creates_sheets(self):
        """Test that format creates expected sheets."""
        from openpyxl import load_workbook

        formatter = ExcelFormatter()
        result = self._create_sample_result()

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            output_path = Path(f.name)

        formatter.format(result, output_path)

        # Load and verify sheets
        wb = load_workbook(output_path)
        sheet_names = wb.sheetnames

        assert "VLLM Operations" in sheet_names
        assert "TensorCast Operations" in sheet_names
        assert "Comparison" in sheet_names
        assert "Summary" in sheet_names

    def test_implements_protocol(self):
        """Test that ExcelFormatter implements FormatterProtocol."""
        formatter = ExcelFormatter()
        assert isinstance(formatter, FormatterProtocol)


class TestFormatterProtocol:
    """Tests for FormatterProtocol."""

    def test_protocol_check(self):
        """Test that ExcelFormatter passes protocol check."""
        assert issubclass(ExcelFormatter, FormatterProtocol)

    def test_custom_formatter_protocol(self):
        """Test creating custom formatter that implements protocol."""

        class CustomFormatter:
            def format(self, result: ComparisonResult, output_path: Path) -> None:
                pass

        assert isinstance(CustomFormatter(), FormatterProtocol)
