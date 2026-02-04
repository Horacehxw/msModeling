"""Excel output formatter for profiling comparison results.

This module provides functionality to generate Excel reports with
multiple sheets containing VLLM operations, TensorCast operations,
comparison data, and summary statistics.
"""

from pathlib import Path
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .base import BaseFormatter, ComparisonResult, OperationMatch


# Color definitions for Excel styling
COLORS = {
    "header_blue": "4472C4",
    "header_green": "2E7D32",
    "header_purple": "7B1FA2",
    "header_orange": "E65100",
    "row_alt_blue": "D9E2F3",
    "row_alt_green": "E8F5E9",
    "row_alt_purple": "E1BEE7",
    "row_alt_orange": "FFE0B2",
    "good_green": "C8E6C9",
    "warn_yellow": "FFF9C4",
    "bad_red": "FFCDD2",
}


class ExcelFormatter(BaseFormatter):
    """Excel output formatter for comparison results.

    Generates an Excel workbook with the following sheets:
    1. VLLM Operations: All VLLM operations with timing data
    2. TensorCast Operations: All TensorCast operations with timing data
    3. Comparison: Side-by-side comparison of matched operations
    4. Summary: Summary statistics and configuration
    """

    def format(self, result: ComparisonResult, output_path: Path) -> None:
        """Generate Excel report from comparison result.

        Args:
            result: Comparison result to format
            output_path: Path to save Excel file
        """
        wb = Workbook()
        wb.remove(wb.active)  # Remove default sheet

        # Create sheets
        self._create_vllm_sheet(wb, result)
        self._create_tensorcast_sheet(wb, result)
        self._create_comparison_sheet(wb, result)
        self._create_summary_sheet(wb, result)

        # Save workbook
        wb.save(output_path)
        print(f"Excel report saved to: {output_path}")

    def _create_vllm_sheet(self, wb: Workbook, result: ComparisonResult) -> None:
        """Create VLLM operations sheet."""
        ws = wb.create_sheet(title="VLLM Operations")

        headers = [
            "Op Type",
            "Count",
            "Total Duration (us)",
            "Avg Duration (us)",
            "Core Type(s)",
            "Input Shapes",
        ]
        col_widths = [30, 10, 18, 16, 20, 50]

        # Write headers
        self._write_headers(ws, headers, col_widths, COLORS["header_blue"])
        ws.freeze_panes = "A2"

        # Sort by total duration
        sorted_ops = sorted(
            result.vllm_operations,
            key=lambda x: x.get("total_duration_us", 0),
            reverse=True,
        )

        for idx, op in enumerate(sorted_ops, 2):
            row_data = [
                op.get("op_type", ""),
                op.get("count", 0),
                op.get("total_duration_us", 0),
                op.get("avg_duration_us", 0),
                ", ".join(op.get("core_types", [])[:2]),
                "; ".join(op.get("input_shapes", [])[:2]),
            ]

            self._write_row(ws, idx, row_data, COLORS["row_alt_blue"], [3, 4])

        # Total row
        total_row = len(sorted_ops) + 2
        self._write_total_row(
            ws,
            total_row,
            ["TOTAL", sum(op.get("count", 0) for op in sorted_ops)],
            sum(op.get("total_duration_us", 0) for op in sorted_ops),
            COLORS["good_green"],
            6,
        )

        self._add_borders(ws)

    def _create_tensorcast_sheet(self, wb: Workbook, result: ComparisonResult) -> None:
        """Create TensorCast operations sheet."""
        ws = wb.create_sheet(title="TensorCast Operations")

        headers = [
            "Op Name",
            "Count",
            "Total Time (us)",
            "Avg Time (us)",
            "Bound",
            "Input Shapes",
        ]
        col_widths = [45, 10, 16, 14, 12, 50]

        self._write_headers(ws, headers, col_widths, COLORS["header_green"])
        ws.freeze_panes = "A2"

        # Sort by total time
        sorted_ops = sorted(
            result.tc_operations,
            key=lambda x: x.get("total_time_us", 0),
            reverse=True,
        )

        for idx, op in enumerate(sorted_ops, 2):
            row_data = [
                op.get("op_name", ""),
                op.get("count", 0),
                op.get("total_time_us", 0),
                op.get("avg_time_us", 0),
                op.get("bound_classification", ""),
                ", ".join(op.get("input_shapes", [])[:3]),
            ]

            self._write_row(ws, idx, row_data, COLORS["row_alt_green"], [3, 4])

        # Total row
        total_row = len(sorted_ops) + 2
        self._write_total_row(
            ws,
            total_row,
            ["TOTAL", sum(op.get("count", 0) for op in sorted_ops)],
            sum(op.get("total_time_us", 0) for op in sorted_ops),
            COLORS["good_green"],
            6,
        )

        self._add_borders(ws)

    def _create_comparison_sheet(self, wb: Workbook, result: ComparisonResult) -> None:
        """Create comparison sheet."""
        ws = wb.create_sheet(title="Comparison")

        headers = [
            "VLLM Op",
            "TC Op(s)",
            "VLLM Duration (us)",
            "TC Duration (us)",
            "Difference (us)",
            "Difference %",
            "Match Status",
        ]
        col_widths = [28, 45, 18, 16, 16, 14, 14]

        self._write_headers(ws, headers, col_widths, COLORS["header_purple"])
        ws.row_dimensions[1].height = 30
        ws.freeze_panes = "A2"

        # Sort by VLLM duration
        sorted_matches = sorted(
            result.matches, key=lambda x: x.vllm_duration_us, reverse=True
        )

        for idx, match in enumerate(sorted_matches, 2):
            tc_ops_str = ", ".join(match.tc_ops) if match.tc_ops else "[Not Matched]"

            row_data = [
                match.vllm_op,
                tc_ops_str,
                match.vllm_duration_us,
                match.tc_duration_us,
                match.difference_us,
                match.difference_pct,
                match.match_status,
            ]

            self._write_comparison_row(ws, idx, row_data, match)

        # Total row
        total_row = len(sorted_matches) + 2
        vllm_total = sum(
            m.vllm_duration_us for m in sorted_matches if m.vllm_op != "[Not Matched]"
        )
        tc_total = sum(m.tc_duration_us for m in sorted_matches if m.tc_ops)

        ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
        ws.cell(row=total_row, column=3, value=vllm_total).number_format = "#,##0.00"
        ws.cell(row=total_row, column=4, value=tc_total).number_format = "#,##0.00"

        total_diff = tc_total - vllm_total
        ws.cell(row=total_row, column=5, value=total_diff).number_format = "#,##0.00"

        if vllm_total > 0:
            total_diff_pct = (total_diff / vllm_total) * 100
            ws.cell(
                row=total_row, column=6, value=total_diff_pct / 100
            ).number_format = "+0.0%;-0.0%;0%"

        for col in range(1, 8):
            ws.cell(row=total_row, column=col).fill = PatternFill(
                start_color="CE93D8", fill_type="solid"
            )
            ws.cell(row=total_row, column=col).font = Font(bold=True)

        self._add_borders(ws)

    def _create_summary_sheet(self, wb: Workbook, result: ComparisonResult) -> None:
        """Create summary statistics sheet."""
        ws = wb.create_sheet(title="Summary")

        # Header
        ws.cell(row=1, column=1, value="Profiling Comparison Summary").font = Font(
            bold=True, size=14
        )
        ws.merge_cells("A1:D1")

        row = 3

        # Configuration section
        ws.cell(row=row, column=1, value="CONFIGURATION").font = Font(bold=True, size=12)
        row += 1

        config = result.config
        config_items = [
            ("Model", config.get("model_id", "N/A")),
            ("Device", config.get("device", "N/A")),
            ("World Size", config.get("world_size", "N/A")),
            ("TP Size", config.get("tp_size", "N/A")),
            ("DP Size", config.get("dp_size", "N/A")),
            ("Expert Parallelism", "Enabled" if config.get("ep") else "Disabled"),
            ("Quantization", config.get("quantize_linear_action", "N/A")),
            ("Batch Size", config.get("num_queries", "N/A")),
            ("Query Length", config.get("query_length", "N/A")),
            ("Context Length", config.get("context_length", "N/A")),
        ]

        for label, value in config_items:
            ws.cell(row=row, column=1, value=label)
            ws.cell(row=row, column=2, value=str(value))
            row += 1

        # VLLM Statistics
        row += 1
        ws.cell(row=row, column=1, value="VLLM PROFILING").font = Font(bold=True, size=12)
        row += 1

        summary = result.summary
        vllm_items = [
            ("Total Duration (us)", f"{summary.vllm_total_time_us:.2f}"),
            ("Total Duration (ms)", f"{summary.vllm_total_time_us / 1000:.3f}"),
            ("Unique Op Types", len(result.vllm_operations)),
        ]

        for label, value in vllm_items:
            ws.cell(row=row, column=1, value=label)
            ws.cell(row=row, column=2, value=value)
            row += 1

        # TensorCast Statistics
        row += 1
        ws.cell(row=row, column=1, value="TENSORCAST SIMULATION").font = Font(
            bold=True, size=12
        )
        row += 1

        tc_items = [
            ("Total Duration (us)", f"{summary.tc_total_time_us:.2f}"),
            ("Total Duration (ms)", f"{summary.tc_total_time_us / 1000:.3f}"),
            ("Unique Op Types", len(result.tc_operations)),
        ]

        for label, value in tc_items:
            ws.cell(row=row, column=1, value=label)
            ws.cell(row=row, column=2, value=value)
            row += 1

        # Comparison Statistics
        row += 1
        ws.cell(row=row, column=1, value="COMPARISON METRICS").font = Font(
            bold=True, size=12
        )
        row += 1

        comparison_items = [
            (
                "Overall Time Difference",
                f"{summary.overall_diff_us:+.2f} us ({summary.overall_diff_pct:+.1f}%)",
            ),
            ("Matched Operations", summary.num_matched),
            ("VLLM-only (missing in TC)", summary.num_missing_in_tc),
            ("TC-only (missing in VLLM)", summary.num_missing_in_vllm),
            ("Time Coverage (matched)", f"{summary.coverage_pct:.1f}%"),
            ("Avg Abs Mismatch (matched)", f"{summary.avg_abs_diff_pct:.1f}%"),
        ]

        for label, value in comparison_items:
            ws.cell(row=row, column=1, value=label)
            ws.cell(row=row, column=2, value=str(value))
            row += 1

        # Accuracy breakdown
        row += 1
        ws.cell(row=row, column=1, value="ACCURACY BREAKDOWN").font = Font(
            bold=True, size=12
        )
        row += 1

        matched = [
            m
            for m in result.matches
            if m.match_status not in ("missing_in_tc", "missing_in_vllm")
        ]
        within_20 = len([m for m in matched if abs(m.difference_pct) <= 20])
        within_50 = len([m for m in matched if 20 < abs(m.difference_pct) <= 50])
        over_50 = len([m for m in matched if abs(m.difference_pct) > 50])

        accuracy_items = [
            ("Within 20% (excellent)", f"{within_20} ops"),
            ("20-50% (acceptable)", f"{within_50} ops"),
            ("Over 50% (needs review)", f"{over_50} ops"),
        ]

        for label, value in accuracy_items:
            ws.cell(row=row, column=1, value=label)
            ws.cell(row=row, column=2, value=value)
            row += 1

        # Column widths
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 50

        self._add_borders(ws)

    def _write_headers(
        self, ws, headers: List[str], widths: List[int], color: str
    ) -> None:
        """Write header row with styling."""
        for col, (header, width) in enumerate(zip(headers, widths), 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color=color, fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[cell.column_letter].width = width

    def _write_row(
        self, ws, row: int, data: List, alt_color: str, numeric_cols: List[int]
    ) -> None:
        """Write data row with alternating colors."""
        for col, val in enumerate(data, 1):
            cell = ws.cell(row=row, column=col, value=val)
            if row % 2 == 0:
                cell.fill = PatternFill(start_color=alt_color, fill_type="solid")
            if col in numeric_cols:
                cell.number_format = "#,##0.00"

    def _write_comparison_row(
        self, ws, row: int, data: List, match: OperationMatch
    ) -> None:
        """Write comparison row with status-based coloring."""
        for col, val in enumerate(data, 1):
            cell = ws.cell(row=row, column=col, value=val)

            if row % 2 == 0:
                cell.fill = PatternFill(
                    start_color=COLORS["row_alt_purple"], fill_type="solid"
                )

            # Format numbers
            if col in (3, 4, 5):
                cell.number_format = "#,##0.00"
            elif col == 6:
                # Color code difference percentage
                if val == float("inf"):
                    cell.value = "N/A"
                    cell.font = Font(color="888888")
                else:
                    cell.number_format = "+0.0%;-0.0%;0%"
                    cell.value = val / 100.0 if val else 0

                    if abs(val) <= 20:
                        cell.font = Font(color="008000")
                        cell.fill = PatternFill(
                            start_color=COLORS["good_green"], fill_type="solid"
                        )
                    elif abs(val) <= 50:
                        cell.font = Font(color="FF6600")
                        cell.fill = PatternFill(
                            start_color=COLORS["warn_yellow"], fill_type="solid"
                        )
                    else:
                        cell.font = Font(color="FF0000", bold=True)
                        cell.fill = PatternFill(
                            start_color=COLORS["bad_red"], fill_type="solid"
                        )

            # Status coloring
            if col == 7:
                if val == "exact":
                    cell.font = Font(color="008000")
                elif val == "partial":
                    cell.font = Font(color="FF6600")
                else:
                    cell.font = Font(color="FF0000")

    def _write_total_row(
        self, ws, row: int, prefix_data: List, total_time: float, color: str, num_cols: int
    ) -> None:
        """Write total row with styling."""
        ws.cell(row=row, column=1, value=prefix_data[0]).font = Font(bold=True)
        if len(prefix_data) > 1:
            ws.cell(row=row, column=2, value=prefix_data[1])
        ws.cell(row=row, column=3, value=total_time).number_format = "#,##0.00"

        for col in range(1, num_cols + 1):
            ws.cell(row=row, column=col).fill = PatternFill(
                start_color=color, fill_type="solid"
            )
            ws.cell(row=row, column=col).font = Font(bold=True)

    def _add_borders(self, ws) -> None:
        """Add borders to all cells with data."""
        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )
        for row in ws.iter_rows(
            min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column
        ):
            for cell in row:
                if cell.value is not None:
                    cell.border = border
